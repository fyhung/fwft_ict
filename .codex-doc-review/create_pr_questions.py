from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = r'DSE PR Qns\DSE PR Questions.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.72); sec.right_margin = Inches(.72)
normal = doc.styles['Normal']
normal.font.name = 'Arial'; normal.font.size = Pt(11)
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'PMingLiU')
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.05

# Worksheet title follows the existing series naming, with restrained styling.
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
r=p.add_run('DSE PR Questions'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(15)

def shade(cell, fill='D9D9D9'):
    pr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); pr.append(shd)
def set_cell(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text=''
    p=cell.paragraphs[0]; p.alignment=align; p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
    r=p.add_run(str(text)); r.bold=bold; r.font.name='Arial'; r.font.size=Pt(10.5); r._element.rPr.rFonts.set(qn('w:eastAsia'),'PMingLiU')
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def grid_table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers): set_cell(t.rows[0].cells[i], h, True); shade(t.rows[0].cells[i])
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row): set_cell(cells[i], x)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def para(text, bold=False, indent=0, before=0, after=5):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(indent); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after)
    r=p.add_run(text); r.bold=bold; r.font.name='Arial'; r.font.size=Pt(11); r._element.rPr.rFonts.set(qn('w:eastAsia'),'PMingLiU')
    return p

para('小明建構了下列鏈表來儲存學生的英文姓名，並以陣列來顯示此鏈表。在此鏈表中，有一指示標 Next 儲存下一個節點的地址。首個節點儲存了「START」。', after=5)
para('地址　　　內容　　　Next', bold=True, indent=.85, after=0)
grid_table(['地址','內容','Next'], [[0,'START',3],[1,'Ben',4],[2,'Kate',-1],[3,'Amy',4],[4,'Jade',5],[5,'Elle',-1]], [0.8,1.4,1.0])
para('(a) (i) 順序寫出「START」後兩個節點的內容。', before=2, after=7)
para('(ii) 小明利用「-1」來表示一個空指示標。除了「-1」外，舉出小明可採用的數值範圍。', after=8)
para('(iii) 包括首個節點「START」，此鏈表共有多少個節點？（4 分）', after=9)

para('小明加建另一個指示標 Previous 而設計了 LL1。在每一個節點內，Previous 指向之前一個節點，如以下例子所示：', after=3)
para('LL1', bold=True, indent=.7, after=2)
grid_table(['地址','內容','Previous','Next'], [[0,'START',-1,3],[1,'John',3,4],[2,'',''],[3,'Susan',0,1],[4,'Fiona',1,-1]], [.7,1.25,1.35,1.05])
para('小明設計了兩個操作 DELETE 及 INSERT。DELETE(n) 會刪除第 n 個節點，而 INSERT(n, sname) 會在第 n 個節點後加入一個節點，其內容為 sname。首個節點儲存了「START」。', after=5)
para('(b) 小明順序執行下列操作來更新以上 LL1 的例子：', after=3)
para('INSERT(4, \'Mary\')\nDELETE(3)', indent=.7, after=5)
para('在下方更新 LL1。', after=2)
para('LL1', bold=True, indent=.7, after=1)
grid_table(['地址','內容','Previous','Next'], [[0,'START',-1,''],[1,'','',''],[2,'',''],[3,'','',''],[4,'','','']], [.7,1.25,1.35,1.05])
para('(4 分)', after=7)
para('(c) 舉出在設計內加進了 Previous 的一個優點和一個缺點。（2 分）', after=8)

para('(d) 小明修改 LL1 而設計了 LL2，他採用 PTR 來取得 Previous 及 Next。每個節點的 PTR 儲存 Previous 及 Next 內地址的總和（即 PTR = Previous + Next）。', after=4)
para('例如：就下列 LL1，', after=2)
para('LL1', bold=True, indent=.7, after=1)
grid_table(['地址','內容','Previous','Next'], [[0,'START',-1,3],[1,'John',3,4],[2,'',''],[3,'Susan',0,1],[4,'Fiona',1,-1],[5,'','']], [.7,1.25,1.35,1.05])
para('其對應的 LL2 會是：', after=2)
para('LL2', bold=True, indent=.7, after=1)
grid_table(['地址','內容','PTR'], [[0,'START',2],[1,'John',7],[2,'',''],[3,'Susan',1],[4,'Fiona',0],[5,'','']], [.8,1.4,1.0])
para('(i) 在下列 LL2，順序寫出「START」後三個節點的內容。', before=1, after=2)
para('LL2', bold=True, indent=.7, after=1)
grid_table(['地址','內容','PTR'], [[0,'START',0],[1,'Candy',3],[2,'Ben',3],[3,'Amy',6],[4,'Lee',7],[5,'Daisy',7]], [.8,1.4,1.0])
para('(ii) 採用 LL2 的好處是什麼？（4 分）', after=0)

doc.core_properties.title='DSE PR Questions'
doc.core_properties.subject='Linked lists and pointer representation questions'
doc.save(out)
print(out)
