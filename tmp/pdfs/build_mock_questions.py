from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r'C:\Users\teacher\github\fwft_ict')
OUT = ROOT / 'DSE PR Qns' / 'DSE 資料結構模擬題.docx'
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin = Cm(1.5), Cm(1.5)
s.left_margin, s.right_margin = Cm(1.8), Cm(1.8)
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(10.5)
normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title', 19), ('Heading 1', 15), ('Heading 2', 11)]:
    st = doc.styles[name]
    st.font.name = 'Arial'; st.font.size = Pt(size); st.font.color.rgb = __import__('docx').shared.RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    st.paragraph_format.space_before = Pt(0); st.paragraph_format.space_after = Pt(7)

def p(text, bold=False, after=5, keep=False):
    a = doc.add_paragraph()
    a.paragraph_format.space_after = Pt(after)
    a.paragraph_format.keep_with_next = keep
    a.add_run(text).bold = bold
    return a

def part(label, text, marks):
    return p(f'({label}) {text}\u00a0（{marks}\u00a0分）')

def code(text):
    a = doc.add_paragraph()
    a.paragraph_format.left_indent = Cm(.45)
    a.paragraph_format.space_after = Pt(6)
    a.paragraph_format.line_spacing = 1.0
    a.paragraph_format.keep_together = True
    r = a.add_run(text); r.font.name = 'Consolas'; r.font.size = Pt(9.5)
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    return a

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for row in [headers] + rows:
        cells = t.rows[0].cells if row is headers else t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text=str(value)
            cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths: cells[i].width=Cm(widths[i])
            for a in cells[i].paragraphs:
                a.alignment=WD_ALIGN_PARAGRAPH.CENTER
                a.paragraph_format.space_after=Pt(2); a.paragraph_format.space_before=Pt(2)
                for r in a.runs: r.font.size=Pt(10); r.bold = row is headers
        pr = cells[0]._tc.getparent().get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
    for c in t.rows[0].cells:
        sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8E8E8'); c._tc.get_or_add_tcPr().append(sh)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    p('',after=1)
    return t

def question(n, title, category):
    if n>1: doc.add_page_break()
    if n==1:
        doc.add_paragraph('DSE 資料結構模擬題', 'Title')
        p('資訊及通訊科技　選修部分 D　軟件開發', bold=True, after=3)
        p('共 8 題，總分 96 分。請回答全部題目。陣列索引及堆疊表示方法以各題的定義為準。偽代碼的計數迴圈每次增加 1；起始值大於終止值時，不執行迴圈。',after=8)
    doc.add_paragraph(f'第 {n} 題　{title}', 'Heading 1')
    p(f'{category}　｜　12 分',bold=True,after=7)

question(1, '活動記錄的還原', '堆疊')
p('某繪圖程式以堆疊 U 儲存已完成的動作，以堆疊 R 儲存可重做的動作。下列操作均作用於堆疊頂部：')
table(['操作','說明'],[['push(S, x)','把 x 放入堆疊 S'],['pop(S)','移除並傳回 S 的頂部元素'],['isEmpty(S)','若 S 是空的，傳回 TRUE；否則傳回 FALSE']],[5,12.4])
p('本題以 [底部, …, 頂部] 表示堆疊。最初 U = [A, B, C]，R = []。',keep=True)
code('push(R, pop(U))\npush(R, pop(U))\npush(U, pop(R))\npush(U, D)')
part('a','寫出上述操作後 U 和 R 的內容。',3)
part('b','使用者完成新動作 D 後，程式應清空 R。解釋清空 R 的原因。',2)
part('c','完成以下子程序 Redo，把 R 頂部的動作移回 U。若 R 是空的，傳回 FALSE；操作成功則傳回 TRUE。',4)
code('Redo\n    如果 ① __________ 則\n        傳回 FALSE\n    否則\n        temp ← ② __________\n        ③ __________\n        傳回 ④ __________')
part('d','U 最多可儲存 50 個動作。若 U 已滿而有新動作，程式希望移除最舊的動作。解釋只使用一個堆疊為何難以直接完成此操作，並建議一個處理方法。',3)

question(2, '利用堆疊計算算式', '堆疊')
p('小晴以堆疊 S 計算後序算式。算式以陣列 A 儲存，每個元素是一個整數或運算符號「+」、「−」。運算符號出現時，先取出的數值是右運算元，再取出的是左運算元。例如 [8, 3, −] 的結果是 5。')
p('push、pop 及 isEmpty 的定義與第 1 題相同。size(S) 傳回 S 內的元素數目。每次計算開始前 S 均為空；本題以 [底部, …, 頂部] 表示堆疊。')
part('a','A = [9, 4, −, 2, +]。寫出處理第 3 個元素和第 5 個元素後 S 的內容，以及算式的最終結果。',3)
part('b','分別指出下列算式是否有效，並解釋原因：\n　(i) [6, +, 2]\n　(ii) [6, 2, 1, +]',3)
part('c','完成下列偽代碼。若算式有效，傳回結果；否則傳回文字「無效」。clear(S) 會清空 S；A 的索引由 1 至 n。',4)
code('Evaluate(A, n)\n    clear(S)\n    設 i 由 1 至 n 執行\n        如果 A[i] 是整數 則\n            push(S, A[i])\n        否則\n            如果 ① __________ 則 傳回「無效」\n            right ← pop(S)\n            left ← pop(S)\n            如果 A[i] =「+」則 push(S, left + right)\n            否則 push(S, ② __________)\n    如果 ③ __________ 則\n        傳回 ④ __________\n    否則 傳回「無效」')
part('d','有人把堆疊 S 改為隊列，其他處理規則不變。以算式 [8, 3, −] 為例，指出所得結果，並解釋兩種數據結構的取出次序如何影響計算。',2)

question(3, '圖書館預約的循環隊列', '隊列')
p('圖書館以循環隊列 Q 儲存預約者的編號。陣列 A 的索引由 0 至 5；H 儲存前端元素的索引，N 儲存元素數目。最多可儲存 6 個元素。deq(Q) 移除並傳回前端元素；移除後把原位置設為「空」。')
table(['索引','0','1','2','3','4','5'],[['A','C','空','空','空','A','B']],[2.4,2.5,2.5,2.5,2.5,2.5,2.5])
p('最初 H = 4，N = 3，邏輯次序為 A、B、C。enq(Q, x) 把 x 加至隊列末端。')
part('a','順序執行 deq(Q)、enq(Q, D)、enq(Q, E)、deq(Q)。寫出最後 A 的各元素、H 和 N。',3)
part('b','當 N > 0 時，寫出隊列末端元素索引的公式，以 H 和 N 表示。',2)
part('c','完成以下子程序 Enqueue。mod 表示取餘數；若隊列已滿，傳回 FALSE。',4)
code('Enqueue(x)\n    如果 ① __________ 則\n        傳回 FALSE\n    否則\n        i ← ② __________\n        A[i] ← x\n        N ← ③ __________\n        傳回 ④ __________')
part('d','當隊列是空的，H 和 N 應符合甚麼條件？解釋為何不能只以 H = 0 判斷隊列是否為空。',2)
part('e','本題另設 N 記錄元素數目。與只用前端及下一個插入位置、並保留一個空位置的循環隊列相比，本題可多儲存多少個元素？',1)

question(4, '列印工作管理', '隊列')
p('某電腦室以隊列 Q 處理列印工作，先到的工作先列印。A 是索引由 0 至 7 的陣列，N 是工作數目。前端固定在 A[0]，有效工作儲存在 A[0] 至 A[N−1]。')
p('最初 N = 4，A 的內容如下。每個編號代表一項完整工作。',keep=True)
table(['索引','0','1','2','3','4','5','6','7'],[['A','P1','P2','P3','P4','空','空','空','空']])
part('a','順序加入 P5、移除前端工作、加入 P6、移除前端工作。寫出最後有效工作的次序和 N。',3)
part('b','完成以下 Dequeue 偽代碼。移除前端工作後，程式把其餘工作向前移動，並清空不再使用的位置。',3)
code('Dequeue\n    如果 N = 0 則 傳回「沒有工作」\n    temp ← A[0]\n    設 i 由 0 至 ① __________ 執行\n        A[i] ← ② __________\n    N ← N − 1\n    A[③ __________] ←「空」\n    傳回 temp')
part('c','程式改用 H 記錄前端，並以循環隊列實作。\n　(i) 說明移除工作時，H 應如何更新。\n　(ii) 比較兩種方法在有 N 個工作時，每次移除工作所需的陣列元素搬移次數。',3)
part('d','電腦室新增「緊急」及「一般」兩類工作。同類工作仍按先到先處理，但緊急工作優先。\n　(i) 建議如何用兩個隊列實作此規則。\n　(ii) 若緊急工作不斷到達，一般工作可能遇到甚麼問題？建議一項公平處理規則。',3)

question(5, '以陣列表示單向鏈表', '鏈表')
p('某學會以單向鏈表儲存會員姓名。首節點在地址 0，內容為 START；Next 儲存下一個節點的地址，−1 表示空指示標。空置節點不屬於鏈表。')
table(['地址','內容','Next'],[[0,'START',4],[1,'Leo',5],[2,'空','空'],[3,'Ivy',-1],[4,'Ada',1],[5,'Mia',3],[6,'空','空']],[3,7,7.4])
part('a','順序寫出 START 後的所有姓名，並指出不包括 START 的節點數目。',3)
p('INSERT(n, name) 在地址 n 的節點後插入新姓名，使用地址最小的空置節點。DELETE(n) 刪除地址 n 的節點，並把該節點的內容及 Next 設為「空」。兩個操作均不容許刪除 START。')
part('b','順序執行 INSERT(1, Noel) 及 DELETE(5)。列出所有被修改節點的地址、內容和 Next，包括被釋放的節點。',4)
part('c','以下偽代碼把地址 new 的新節點插入地址 n 的節點之後。已知 n 有效而 new 是空置節點。完成偽代碼，並解釋為何兩個指示標更新步驟的先後次序十分重要。',3)
code('InsertAfter(n, new, name)\n    Content[new] ← name\n    Next[new] ← ① __________\n    Next[n] ← ② __________')
part('d','若改用連續的陣列、按姓名出現次序儲存資料，比較在中間插入會員時，陣列與鏈表各自需要修改甚麼。',2)

question(6, '雙向鏈表與指示標壓縮', '鏈表')
p('音樂播放器以雙向鏈表儲存歌單。地址 0 的 START 為首節點。Previous 及 Next 分別儲存前一個和下一個節點的地址；−1 表示空指示標。')
table(['地址','內容','Previous','Next'],[[0,'START',-1,3],[1,'Blue',3,4],[2,'空','空','空'],[3,'Rain',0,1],[4,'Wind',1,-1],[5,'空','空','空']],[2.4,6,4.5,4.5])
part('a','寫出歌單的播放次序；若目前播放地址 4 的歌曲，寫出連續向前移動兩次所到達的節點內容。',3)
part('b','刪除地址 1 的歌曲。寫出需要修改的兩個相鄰節點及其新指示標值，並把地址 1 的內容及兩個指示標設為「空」。',4)
p('播放器考慮把每個有效節點的兩個指示標改為一個整數 PTR，其中 PTR = Previous + Next。START 的前一個地址仍視為 −1。',keep=True)
part('c','本小題使用上方原來的歌單，並不使用 (b) 的結果。\n　(i) 計算 START 及 Rain 的 PTR。\n　(ii) 已知目前節點的地址、PTR 及前一個節點的地址 prev，寫出下一個節點地址的公式。',3)
part('d','相對於單向鏈表，舉出雙向鏈表的一個優點和一個缺點。',2)

question(7, '調整派發次序', '混合　隊列與堆疊')
p('老師以隊列 Q 儲存待派發的教材編號，並可使用輔助堆疊 S。enq(Q, x) 把 x 加至隊列末端；deq(Q) 移除並傳回前端元素。push 及 pop 的定義與第 1 題相同。')
p('隊列以 [前端, …, 末端] 表示；堆疊以 [底部, …, 頂部] 表示。最初 Q = [11, 22, 33, 44]，S 是空的。',keep=True)
code('push(S, deq(Q))\npush(S, deq(Q))\nenq(Q, pop(S))\nenq(Q, pop(S))')
part('a','寫出上述操作後 Q 和 S 的內容。',3)
part('b','完成以下 ReverseAll，使 Q 內所有元素的次序倒轉。isEmpty(X) 可檢查隊列或堆疊是否為空；S 在開始時是空的。',4)
code('ReverseAll\n    當 ① __________ 執行\n        push(S, ② __________)\n    當 ③ __________ 執行\n        enq(Q, ④ __________)')
part('c','編寫 ReverseFirst(Q, k, n) 的偽代碼，只倒轉 Q 的前 k 個元素，並保留其餘元素的相對次序；n 是 Q 的原有元素數目，1 ≤ k ≤ n。只可使用一個空堆疊 S、隊列操作和簡單整數變量。\n例如：[11, 22, 33, 44, 55] 在 k = 3 時應變成 [33, 22, 11, 44, 55]。',3)
part('d','在 (b) 中，若把輔助堆疊改為另一個空隊列，並把 push、pop 分別改為 enq、deq，所得次序有何變化？解釋原因。',2)

question(8, '以鏈表實作客服隊列', '混合　隊列與鏈表')
p('客服系統按先到先處理的次序服務客戶。系統以單向鏈表實作隊列，每個節點包含 ID 和 Next；H、T 分別儲存前端及末端節點的地址。−1 是空指示標。')
table(['地址','ID','Next'],[[2,'C03',-1],[5,'C01',7],[7,'C02',2]],[3,7,7.4])
p('最初 H = 5，T = 2。allocate() 配置新節點並傳回其地址；release(x) 釋放地址 x 的節點。')
part('a','說明本系統應採用隊列而不是堆疊的原因，並寫出最初的服務次序。',2)
part('b','新客戶 C04 加入，allocate() 傳回地址 9。其後服務並移除前端客戶。列出這兩個操作後 H、T，以及地址 2 和 9 的 ID、Next。',3)
part('c','完成以下兩段偽代碼。新節點配置一定成功；Dequeue 只會在隊列非空時呼叫。',4)
code('Enqueue(id)\n    new ← allocate()\n    ID[new] ← id\n    Next[new] ← −1\n    如果 H = −1 則\n        H ← new\n    否則\n        ① __________ ← new\n    T ← ② __________\n\nDequeue\n    old ← H\n    value ← ID[old]\n    H ← ③ __________\n    如果 H = −1 則\n        T ← ④ __________\n    release(old)\n    傳回 value')
part('d','與使用固定大小陣列的循環隊列相比：\n　(i) 舉出鏈表隊列在容量方面的一個優點。\n　(ii) 舉出其記憶體使用上的一個缺點。\n　(iii) 另一子程序以 T = −1 判斷隊列是否為空。若移除唯一的節點後沒有把 T 設為 −1，此判斷會出現甚麼錯誤？',3)

# Footer page field, kept minimal for classroom printing.
footer = s.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('資料結構模擬題　｜　')
r.font.size = Pt(9)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title='DSE 資料結構模擬題'
doc.core_properties.subject='堆疊、隊列、鏈表及混合應用原創練習'
doc.core_properties.author=''
# Remove the default template's decorative paragraph borders.
for root in [doc._element, doc.styles.element]:
    for border in list(root.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(OUT)
print(OUT)
